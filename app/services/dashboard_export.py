"""Generadores en memoria para exportar el Dashboard sin almacenar archivos."""
import csv
from datetime import datetime
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


HEADERS_REPORTS = [
    "ID", "Fecha", "Hora", "Técnico / trabajador", "Cliente", "Área", "Sección", "Maquinaria",
    "Tipo de servicio", "Descripción", "Comienzo de tarea", "Finalización de tarea",
]
HEADERS_FUEL = [
    "ID", "Fecha", "Hora", "Técnico / trabajador", "Litros Generador 1", "Horómetro Generador 1",
    "Litros Generador 2", "Horómetro Generador 2", "Observaciones",
]
HEADERS_CSV = ["Tipo de registro", *HEADERS_REPORTS, "Litros Generador 1", "Horómetro Generador 1", "Litros Generador 2", "Horómetro Generador 2", "Observaciones"]


def _naive(value):
    """Excel no admite datetimes con zona, pero sí valores fecha/hora nativos."""
    if isinstance(value, datetime) and value.tzinfo:
        return value.replace(tzinfo=None)
    return value


def _date(value):
    return _naive(value).date() if isinstance(value, datetime) else None


def _time(value):
    return _naive(value).time().replace(tzinfo=None) if isinstance(value, datetime) else None


def report_row(report):
    return [
        report.id, _date(report.created_at), _time(report.created_at), report.technician.full_name,
        report.client, report.area.name if report.area else None, report.section.name if report.section else None,
        report.machinery.name if report.machinery else None, report.service_type, report.description,
        _naive(report.task_started_at), _naive(report.task_finished_at),
    ]


def fuel_row(load, technician, detail_map):
    generator_1 = detail_map.get((load.id, 1))
    generator_2 = detail_map.get((load.id, 2))
    return [
        load.id, _date(load.loaded_at), _time(load.loaded_at), technician.full_name if technician else "Sin registro",
        float(generator_1.liters) if generator_1 else None, float(generator_1.hourmeter) if generator_1 else None,
        float(generator_2.liters) if generator_2 else None, float(generator_2.hourmeter) if generator_2 else None,
        load.observations,
    ]


def _style_sheet(sheet, widths):
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "dd-mm-yyyy hh:mm"
            elif getattr(cell.value, "year", None) and getattr(cell.value, "month", None):
                cell.number_format = "dd-mm-yyyy"


def xlsx_bytes(summary_data, reports, fuel_loads, filter_labels):
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Resumen"
    summary_sheet.append(["Dashboard de Reportes Mantención"])
    summary_sheet["A1"].font = Font(bold=True, size=14, color="1D4ED8")
    summary_sheet.append([])
    summary_sheet.append(["Filtros aplicados", "Valor"])
    for label, value in filter_labels:
        summary_sheet.append([label, value or "Todos"])
    summary_sheet.append([])
    summary_sheet.append(["Indicador", "Cantidad"])
    for label, value in (
        ("Total de registros", summary_data["summary"]["total_records"]),
        ("Total de mantenciones", summary_data["summary"]["maintenance_reports"]),
        ("Total de cargas de petróleo", summary_data["summary"]["fuel_loads"]),
        ("Días con actividad", summary_data["summary"]["activity_days"]),
    ):
        summary_sheet.append([label, value])
    for row in (3, 13):
        for cell in summary_sheet[row]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 38

    reports_sheet = workbook.create_sheet("Reportes")
    reports_sheet.append(HEADERS_REPORTS)
    for report in reports:
        reports_sheet.append(report_row(report))
    _style_sheet(reports_sheet, [10, 13, 11, 28, 20, 20, 20, 28, 28, 50, 20, 20])

    fuel_sheet = workbook.create_sheet("Cargas de petróleo")
    fuel_sheet.append(HEADERS_FUEL)
    for load, technician, detail_map in fuel_loads:
        fuel_sheet.append(fuel_row(load, technician, detail_map))
    _style_sheet(fuel_sheet, [10, 13, 11, 28, 20, 22, 20, 22, 50])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def csv_bytes(reports, fuel_loads):
    """CSV plano UTF-8 BOM para que Excel interprete acentos y ñ correctamente."""
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(HEADERS_CSV)
    for report in reports:
        row = report_row(report)
        writer.writerow(["Mantención", *row, "", "", "", "", ""])
    for load, technician, detail_map in fuel_loads:
        fuel = fuel_row(load, technician, detail_map)
        writer.writerow(["Carga de petróleo", fuel[0], fuel[1], fuel[2], fuel[3], "", "", "", "", "Carga de petróleo", "", "", "", fuel[4], fuel[5], fuel[6], fuel[7], fuel[8]])
    return ("\ufeff" + output.getvalue()).encode("utf-8")
