from datetime import datetime, timezone
import io
import json
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from werkzeug.security import generate_password_hash

from app import create_app, db
from app.models import Area, Machinery, Section, User


class TestConfig:
    TESTING = True
    SECRET_KEY = "web-test-secret"
    API_TOKEN_SECRET = "api-test-secret-at-least-32-characters"
    API_TOKEN_ISSUER = "reportes-test-api"
    API_ACCESS_TOKEN_MINUTES = 15
    API_REFRESH_TOKEN_DAYS = 30
    SQLALCHEMY_DATABASE_URI = "sqlite://"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    SQLALCHEMY_ENGINE_OPTIONS = {}
    UPLOAD_FOLDER = Path("tests/.uploads")
    FILE_STORAGE_BACKEND = "local"
    AZURE_STORAGE_CONTAINER = "reportes-adjuntos"
    MAX_CONTENT_LENGTH = 1024 * 1024
    ADMIN_USERNAME = "unused@example.test"
    ADMIN_PASSWORD = None


@pytest.fixture()
def client():
    app = create_app(TestConfig)
    with app.app_context():
        # SQLite necesita adjuntar explícitamente el esquema dbo usado por SQL Server.
        db.session.execute(text("ATTACH DATABASE ':memory:' AS dbo"))
        raw_connection = db.session.connection().connection.driver_connection
        raw_connection.create_function("getdate", 0, lambda: "2026-09-14 10:00:00")
        raw_connection.create_function("sysdatetimeoffset", 0, lambda: "2026-09-14 10:00:00+00:00")
        db.create_all()
        db.session.add_all([
            User(username="admin@example.test", full_name="Admin", password_hash=generate_password_hash("CorrectHorseBattery1"), role="Administrador", is_active=True, created_at=datetime.now(timezone.utc)),
            User(username="tech@example.test", full_name="Tech", password_hash=generate_password_hash("CorrectHorseBattery1"), role="Técnico/Operativo", is_active=True, created_at=datetime.now(timezone.utc)),
        ])
        area = Area(name="Malta", is_active=True)
        db.session.add(area)
        db.session.flush()
        section = Section(area_id=area.id, name="Despacho", is_active=True)
        db.session.add(section)
        db.session.flush()
        db.session.add(Machinery(section_id=section.id, source_code=1, name="Soplador 1", is_active=True))
        db.session.commit()
        yield app.test_client()
        db.session.remove()
        db.drop_all()
        shutil.rmtree(TestConfig.UPLOAD_FOLDER, ignore_errors=True)


def login(client, username="admin@example.test"):
    response = client.post("/api/v1/auth/login", json={"username": username, "password": "CorrectHorseBattery1"})
    assert response.status_code == 200
    return response.get_json()["data"]


def test_login_me_refresh_rotation_and_logout(client):
    tokens = login(client)
    assert tokens["token_type"] == "Bearer"
    assert "CorrectHorse" not in str(tokens)

    me = client.get("/api/v1/auth/me", headers={"Authorization": f"Bearer {tokens['access_token']}"})
    assert me.status_code == 200
    assert me.get_json()["data"]["role"] == "Administrador"

    refreshed = client.post("/api/v1/auth/refresh", json={"refresh_token": tokens["refresh_token"]})
    assert refreshed.status_code == 200
    new_tokens = refreshed.get_json()["data"]
    assert new_tokens["refresh_token"] != tokens["refresh_token"]

    reused = client.post("/api/v1/auth/refresh", json={"refresh_token": tokens["refresh_token"]})
    assert reused.status_code == 401
    assert reused.get_json()["error"]["code"] == "invalid_refresh_token"

    logout = client.post("/api/v1/auth/logout", json={"refresh_token": new_tokens["refresh_token"]})
    assert logout.status_code == 200
    revoked = client.post("/api/v1/auth/refresh", json={"refresh_token": new_tokens["refresh_token"]})
    assert revoked.status_code == 401


def test_errors_and_roles_use_the_uniform_envelope(client):
    bad_content_type = client.post("/api/v1/auth/login", data="{}")
    assert bad_content_type.status_code == 415
    assert bad_content_type.get_json()["error"]["code"] == "unsupported_media_type"

    invalid = client.post("/api/v1/auth/login", json={"username": "admin@example.test", "password": "wrong"})
    assert invalid.status_code == 401
    assert invalid.get_json()["error"]["request_id"]

    denied = client.get("/api/v1/admin/status", headers={"Authorization": f"Bearer {login(client, 'tech@example.test')['access_token']}"})
    assert denied.status_code == 403
    assert denied.get_json()["error"]["code"] == "insufficient_role"

    allowed = client.get("/api/v1/admin/status", headers={"Authorization": f"Bearer {login(client)['access_token']}"})
    assert allowed.status_code == 200

    not_found = client.get("/api/v1/no-existe")
    assert not_found.status_code == 404
    assert not_found.get_json()["error"]["code"] == "not_found"

    wrong_method = client.get("/api/v1/auth/login")
    assert wrong_method.status_code == 405
    assert wrong_method.get_json()["error"]["code"] == "method_not_allowed"


def checklist(componentes="No", zona="No", obstrucciones="Si"):
    return [
        {"key": "libre_obstrucciones", "answer": obstrucciones},
        {"key": "componentes_mal_estado", "answer": componentes},
        {"key": "falla_constante", "answer": "No aplica"},
        {"key": "equipo_energizado", "answer": "Si"},
        {"key": "zona_limpia", "answer": zona},
        {"key": "falla_solucionada", "answer": "Si"},
        {"key": "protecciones_instaladas", "answer": "No aplica"},
    ]


def report_payload(**overrides):
    data = {
        "client": "Maltexco", "service_type": "Preventiva", "description": "Prueba automática",
        "area_id": 1, "section_id": 1, "machinery_id": 1,
        "task_started_at": "2026-09-14T08:00:00-03:00", "task_finished_at": "2026-09-14T09:00:00-03:00",
        "checklist": checklist(),
    }
    data.update(overrides)
    return data


def auth_headers(client, username="admin@example.test"):
    return {"Authorization": f"Bearer {login(client, username)['access_token']}"}


def test_catalogs_report_history_filters_and_scope(client):
    headers = auth_headers(client)
    areas = client.get("/api/v1/areas", headers=headers)
    assert areas.status_code == 200
    assert areas.get_json()["data"] == [{"id": 1, "name": "Malta"}]
    assert client.get("/api/v1/sections", headers=headers).status_code == 400
    assert client.get("/api/v1/sections?area_id=1", headers=headers).get_json()["data"][0]["name"] == "Despacho"
    assert client.get("/api/v1/machineries?section_id=1", headers=headers).get_json()["data"][0]["name"] == "Soplador 1"

    created = client.post("/api/v1/reports", json=report_payload(), headers=headers)
    assert created.status_code == 201
    report = created.get_json()["data"]
    assert len(report["checklist"]) == 7
    assert report["technician"]["username"] == "admin@example.test"

    listing = client.get("/api/v1/reports?area_id=1&service_type=Preventiva&page=1&page_size=1", headers=headers)
    assert listing.status_code == 200
    assert listing.get_json()["pagination"]["total"] == 1
    detail = client.get(f"/api/v1/reports/{report['id']}", headers=headers)
    assert detail.status_code == 200
    assert detail.get_json()["data"]["id"] == report["id"]

    technician_listing = client.get("/api/v1/reports", headers=auth_headers(client, "tech@example.test"))
    assert technician_listing.status_code == 200
    assert technician_listing.get_json()["pagination"]["total"] == 0
    forbidden_filter = client.get("/api/v1/reports?technician_id=1", headers=auth_headers(client, "tech@example.test"))
    assert forbidden_filter.status_code == 403


def test_multipart_checklist_evidence_and_fuel_loads(client):
    headers = auth_headers(client)
    missing = client.post("/api/v1/reports", json=report_payload(checklist=checklist(componentes="Si")), headers=headers)
    assert missing.status_code == 422
    assert missing.get_json()["error"]["code"] == "evidence_required"

    multipart_report = report_payload(checklist=checklist(componentes="Si"))
    created = client.post(
        "/api/v1/reports", headers=headers,
        data={"payload": json.dumps(multipart_report), "evidence_componentes_mal_estado": (io.BytesIO(b"png"), "evidencia.png")},
        content_type="multipart/form-data",
    )
    assert created.status_code == 201
    checklist_rows = created.get_json()["data"]["checklist"]
    assert next(item for item in checklist_rows if item["key"] == "componentes_mal_estado")["evidence_attachment_id"]
    attachment = created.get_json()["data"]["attachments"][0]
    download = client.get(attachment["download_path"], headers=headers)
    assert download.status_code == 200
    assert download.data == b"png"
    assert client.get(attachment["download_path"], headers=auth_headers(client, "tech@example.test")).status_code == 404

    fuel_payload = {
        "loaded_at": "2026-09-14T10:30:00-03:00", "observations": "Prueba",
        "generators": [{"number": 1, "liters": 749, "hourmeter": 200}, {"number": 2, "liters": 20.5, "hourmeter": 300}],
    }
    images = {"payload": json.dumps(fuel_payload)}
    for name in ("water_1", "oil_1", "water_2", "oil_2"):
        images[name] = (io.BytesIO(b"png"), f"{name}.png")
    fuel = client.post("/api/v1/fuel-loads", headers=headers, data=images, content_type="multipart/form-data")
    assert fuel.status_code == 201
    assert fuel.get_json()["data"]["generators"][0]["liters"] == 749.0
    assert client.get("/api/v1/fuel-loads?page=1&page_size=1", headers=headers).get_json()["pagination"]["total"] == 1
    assert client.get(f"/api/v1/fuel-loads/{fuel.get_json()['data']['id']}", headers=headers).status_code == 200
    image = client.get(f"/api/v1/fuel-loads/{fuel.get_json()['data']['id']}/generators/1/images/water", headers=headers)
    assert image.status_code == 200
    assert image.data == b"png"

    invalid_fuel = dict(fuel_payload)
    invalid_fuel["generators"] = [{"number": 1, "liters": 750, "hourmeter": 1}, {"number": 2, "liters": 1, "hourmeter": 1}]
    assert client.post("/api/v1/fuel-loads", json=invalid_fuel, headers=headers).status_code == 422


@pytest.mark.parametrize("service_type", ["Correctiva", "Preventiva", "Predictiva", "Nueva instalación", "Otro"])
def test_new_reports_accept_only_the_five_maintenance_types(client, service_type):
    headers = auth_headers(client)
    response = client.post("/api/v1/reports", json=report_payload(service_type=service_type), headers=headers)
    assert response.status_code == 201
    assert response.get_json()["data"]["service_type"] == service_type
    dashboard = client.get("/api/v1/dashboard/summary", query_string={"record_type": "maintenance", "service_type": service_type}, headers=headers)
    assert dashboard.status_code == 200
    assert dashboard.get_json()["data"]["summary"]["maintenance_reports"] == 1


def test_new_reports_and_dashboard_filters_reject_unknown_maintenance_type(client):
    headers = auth_headers(client)
    invalid = client.post("/api/v1/reports", json=report_payload(service_type="Inspección"), headers=headers)
    assert invalid.status_code == 422
    assert invalid.get_json()["error"]["code"] == "validation_error"
    assert client.get("/api/v1/dashboard/summary?service_type=Inspecci%C3%B3n", headers=headers).status_code == 400


def test_dashboard_summary_aggregates_in_server_and_enforces_technician_scope(client):
    admin_headers = auth_headers(client)
    created_report = client.post("/api/v1/reports", json=report_payload(), headers=admin_headers)
    assert created_report.status_code == 201
    fuel_payload = {
        "loaded_at": "2026-09-14T10:30:00-03:00", "observations": "Carga de prueba",
        "generators": [{"number": 1, "liters": 10, "hourmeter": 1}, {"number": 2, "liters": 20, "hourmeter": 2}],
    }
    images = {"payload": json.dumps(fuel_payload)}
    for name in ("water_1", "oil_1", "water_2", "oil_2"):
        images[name] = (io.BytesIO(b"png"), f"{name}.png")
    assert client.post("/api/v1/fuel-loads", headers=admin_headers, data=images, content_type="multipart/form-data").status_code == 201

    summary = client.get("/api/v1/dashboard/summary?record_type=all&activity_page=1&activity_page_size=10", headers=admin_headers)
    assert summary.status_code == 200
    data = summary.get_json()["data"]
    assert data["summary"] == {"total_records": 2, "maintenance_reports": 1, "fuel_loads": 1, "activity_days": 1}
    assert {item["record_type"] for item in data["activity"]} == {"maintenance", "fuel"}
    assert data["charts"]["areas"] == [{"id": 1, "name": "Malta", "count": 1}]
    assert {item["id"] for item in data["technicians"]} == {1, 2}

    tech_headers = auth_headers(client, "tech@example.test")
    forbidden = client.get("/api/v1/dashboard/summary?technician_id=1", headers=tech_headers)
    assert forbidden.status_code == 403
    assert forbidden.get_json()["error"]["code"] == "insufficient_role"
    own = client.get("/api/v1/dashboard/summary?technician_id=2", headers=tech_headers)
    assert own.status_code == 200
    assert own.get_json()["data"]["summary"]["total_records"] == 0


def test_dashboard_web_exports_share_filters_format_and_permissions(client):
    """Excel/CSV reutilizan el Dashboard y nunca amplían el alcance del técnico."""
    admin_headers = auth_headers(client)
    created = client.post(
        "/api/v1/reports", json=report_payload(service_type="Predictiva", description="Inspección de ñandú"), headers=admin_headers
    )
    assert created.status_code == 201
    fuel_payload = {
        "loaded_at": "2026-09-14T10:30:00-03:00", "observations": "Revisión con á, é y ñ",
        "generators": [{"number": 1, "liters": 10.5, "hourmeter": 100}, {"number": 2, "liters": 20, "hourmeter": 200}],
    }
    images = {"payload": json.dumps(fuel_payload)}
    for name in ("water_1", "oil_1", "water_2", "oil_2"):
        images[name] = (io.BytesIO(b"png"), f"{name}.png")
    assert client.post("/api/v1/fuel-loads", headers=admin_headers, data=images, content_type="multipart/form-data").status_code == 201

    with client.session_transaction() as web_session:
        web_session.update({"user_id": 1, "user_name": "Admin", "role": "Administrador"})
    filters = "record_type=maintenance&technician_id=1&area_id=1&section_id=1&machinery_id=1&service_type=Predictiva&date_from=2026-09-14&date_to=2026-09-14"
    dashboard_page = client.get(f"/dashboard?{filters}")
    assert dashboard_page.status_code == 200
    assert b"Exportar" in dashboard_page.data
    assert b"dashboard/export" in dashboard_page.data
    xlsx = client.get(f"/dashboard/export?format=xlsx&{filters}")
    assert xlsx.status_code == 200
    assert xlsx.headers["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment;" in xlsx.headers["Content-Disposition"]
    workbook = load_workbook(io.BytesIO(xlsx.data), data_only=True)
    assert workbook.sheetnames == ["Resumen", "Reportes", "Cargas de petróleo"]
    assert workbook["Resumen"]["B14"].value == 1  # Total de registros con los filtros combinados.
    assert workbook["Reportes"].max_row == 2
    assert workbook["Reportes"]["I2"].value == "Predictiva"
    assert "ñandú" in workbook["Reportes"]["J2"].value
    assert workbook["Cargas de petróleo"].max_row == 1

    all_xlsx = load_workbook(io.BytesIO(client.get("/dashboard/export?format=xlsx&record_type=all&technician_id=1").data), data_only=True)
    assert all_xlsx["Cargas de petróleo"].max_row == 2
    assert isinstance(all_xlsx["Cargas de petróleo"]["E2"].value, float)
    assert all_xlsx["Cargas de petróleo"]["I2"].value == "Revisión con á, é y ñ"

    csv_response = client.get("/dashboard/export?format=csv&record_type=all&technician_id=1")
    assert csv_response.status_code == 200
    assert csv_response.headers["Content-Type"].startswith("text/csv")
    assert "attachment;" in csv_response.headers["Content-Disposition"]
    assert csv_response.data.startswith(b"\xef\xbb\xbf")
    assert "Inspección de ñandú" in csv_response.data.decode("utf-8-sig")
    assert "Carga de petróleo" in csv_response.data.decode("utf-8-sig")

    empty = client.get("/dashboard/export?format=xlsx&record_type=maintenance&technician_id=2")
    assert empty.status_code == 200
    assert load_workbook(io.BytesIO(empty.data))["Reportes"].max_row == 1

    with client.session_transaction() as web_session:
        web_session.update({"user_id": 2, "user_name": "Tech", "role": "Técnico/Operativo"})
    forbidden = client.get("/dashboard/export?format=csv&technician_id=1")
    assert forbidden.status_code == 403
    own = client.get("/dashboard/export?format=csv&technician_id=2")
    assert own.status_code == 200


def test_existing_web_forms_continue_to_use_the_shared_rules(client):
    client.get("/auth/login")
    with client.session_transaction() as session:
        csrf = session["csrf_token"]
    assert client.post("/auth/login", data={"username": "admin@example.test", "password": "CorrectHorseBattery1", "csrf_token": csrf}).status_code == 302

    form_page = client.get("/reports/maintenance/new")
    assert form_page.status_code == 200
    for service_type in ("Correctiva", "Preventiva", "Predictiva", "Nueva instalación", "Otro"):
        assert service_type.encode("utf-8") in form_page.data
    assert b"Mantenimiento Preventivo" not in form_page.data
    with client.session_transaction() as session:
        csrf = session["csrf_token"]
    report_form = {
        "csrf_token": csrf, "client": "Maltexco", "service_type": "Nueva instalación", "description": "Formulario web",
        "area_id": "1", "section_id": "1", "machinery_id": "1",
        "task_started_at": "2026-09-14T08:00", "task_finished_at": "2026-09-14T09:00",
    }
    for item in checklist():
        report_form[f"check_{item['key']}"] = item["answer"]
    web_report = client.post("/reports/maintenance/new", data=report_form, follow_redirects=False)
    assert web_report.status_code == 302
    assert "/reports/" in web_report.headers["Location"]

    client.get("/reports/fuel/new")
    with client.session_transaction() as session:
        csrf = session["csrf_token"]
    fuel_form = {"csrf_token": csrf, "loaded_at": "2026-09-14T10:00", "liters_1": "1", "hourmeter_1": "1", "liters_2": "2", "hourmeter_2": "2"}
    for name in ("water_1", "oil_1", "water_2", "oil_2"):
        fuel_form[name] = (io.BytesIO(b"png"), f"{name}.png")
    assert client.post("/reports/fuel/new", data=fuel_form, content_type="multipart/form-data", follow_redirects=False).status_code == 302
